#!/usr/bin/env python3
"""Markdownレポートを配布用のPDF/docxに変換する再利用スクリプト。

使い方:
    python3 scripts/md_report_export.py <input.md> [--pdf] [--docx]
    （フラグ省略時は両方生成。出力は入力と同じディレクトリに同名で .pdf / .docx）

依存: python-markdown, weasyprint (PDF), python-docx (docx)
日本語フォントはmacOSのシステムフォント(Hiragino)を利用。
"""
import sys
import re
from pathlib import Path

PDF_CSS = """
@page { size: A4; margin: 20mm 18mm; }
body { font-family: "Hiragino Sans", "Hiragino Kaku Gothic ProN", "Noto Sans JP", sans-serif;
       font-size: 10.5pt; line-height: 1.75; color: #1a1a1a; }
h1 { font-size: 20pt; color: #0b3d2e; border-bottom: 3px solid #0b3d2e; padding-bottom: 6px; margin-top: 0; }
h2 { font-size: 14pt; color: #0b3d2e; margin-top: 22px; border-left: 5px solid #1a9e6e; padding-left: 10px; }
h3 { font-size: 12pt; color: #13543f; margin-top: 16px; }
p { margin: 8px 0; }
strong { color: #0b3d2e; }
code { font-family: "SF Mono", Menlo, monospace; background: #f1f3f2; padding: 1px 5px; border-radius: 4px; font-size: 9.5pt; }
blockquote { background: #eefaf4; border-left: 5px solid #1a9e6e; margin: 14px 0; padding: 10px 16px; border-radius: 6px; font-weight: 600; }
table { border-collapse: collapse; width: 100%; margin: 12px 0; font-size: 9.5pt; }
th { background: #0b3d2e; color: #fff; text-align: left; padding: 7px 10px; }
td { border: 1px solid #d5ddd9; padding: 6px 10px; vertical-align: top; }
tr:nth-child(even) td { background: #f7faf9; }
hr { border: none; border-top: 1px solid #d5ddd9; margin: 20px 0; }
ul, ol { margin: 8px 0; padding-left: 22px; }
li { margin: 4px 0; }
em { color: #555; font-style: normal; }
"""


def to_pdf(md_text: str, out: Path):
    import markdown
    from weasyprint import HTML
    body = markdown.markdown(md_text, extensions=["tables", "fenced_code", "sane_lists"])
    html = f"<html><head><meta charset='utf-8'><style>{PDF_CSS}</style></head><body>{body}</body></html>"
    HTML(string=html).write_pdf(str(out))
    print(f"  PDF : {out}")


# ---- docx: 対象レポートで使う構文に絞った軽量パーサ ----
INLINE = re.compile(r"(\*\*.+?\*\*|`.+?`)")


def _add_inline(paragraph, text):
    """**bold** と `code` を run に分解して段落へ追加。"""
    for part in INLINE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2]); run.bold = True
        elif part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1]); run.font.name = "Menlo"
        else:
            paragraph.add_run(part)


def to_docx(md_text: str, out: Path):
    from docx import Document
    from docx.shared import Pt, RGBColor

    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Hiragino Sans"
    style.font.size = Pt(10.5)

    lines = md_text.split("\n")
    i = 0
    while i < len(lines):
        line = lines[i].rstrip()
        # table block
        if line.startswith("|") and i + 1 < len(lines) and re.match(r"^\|[\s:|-]+\|?\s*$", lines[i + 1]):
            header = [c.strip() for c in line.strip("|").split("|")]
            rows = []
            j = i + 2
            while j < len(lines) and lines[j].strip().startswith("|"):
                rows.append([c.strip() for c in lines[j].strip().strip("|").split("|")])
                j += 1
            t = doc.add_table(rows=1, cols=len(header)); t.style = "Light Grid Accent 1"
            for k, h in enumerate(header):
                cell = t.rows[0].cells[k]
                cell.paragraphs[0].clear()
                _add_inline(cell.paragraphs[0], h)
                for r in cell.paragraphs[0].runs:
                    r.bold = True
            for row in rows:
                cells = t.add_row().cells
                for k, val in enumerate(row[: len(header)]):
                    cells[k].paragraphs[0].clear()
                    _add_inline(cells[k].paragraphs[0], val)
            doc.add_paragraph()
            i = j
            continue
        if line.startswith("# "):
            h = doc.add_heading(level=0); _add_inline(h, line[2:])
        elif line.startswith("## "):
            h = doc.add_heading(level=1); _add_inline(h, line[3:])
        elif line.startswith("### "):
            h = doc.add_heading(level=2); _add_inline(h, line[4:])
        elif line.startswith("> "):
            p = doc.add_paragraph(); p.style = "Intense Quote"; _add_inline(p, line[2:])
        elif re.match(r"^[-*] ", line):
            p = doc.add_paragraph(style="List Bullet"); _add_inline(p, line[2:])
        elif re.match(r"^\d+\. ", line):
            p = doc.add_paragraph(style="List Number"); _add_inline(p, re.sub(r"^\d+\.\s*", "", line))
        elif line.strip() == "---":
            doc.add_paragraph("_" * 40)
        elif line.strip() == "":
            pass
        else:
            p = doc.add_paragraph(); _add_inline(p, line)
        i += 1
    doc.save(str(out))
    print(f"  DOCX: {out}")


def main():
    args = sys.argv[1:]
    if not args:
        print(__doc__); sys.exit(1)
    src = Path(args[0])
    want_pdf = "--pdf" in args or ("--docx" not in args)
    want_docx = "--docx" in args or ("--pdf" not in args)
    md = src.read_text(encoding="utf-8")
    print(f"converting {src.name}:")
    if want_pdf:
        to_pdf(md, src.with_suffix(".pdf"))
    if want_docx:
        to_docx(md, src.with_suffix(".docx"))


if __name__ == "__main__":
    main()
